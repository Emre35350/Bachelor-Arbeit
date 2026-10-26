"""Regression tests for the retrieval metrics (app/metrics.py).

Focus on the retrieval-quality trio — Recall@k / Precision@k / F1 — and the rules
around gold (unanswerable controls excluded from those means).

Run:  python -m tests.test_metrics      (no pytest needed)
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import metrics  # noqa: E402


def check(label, got, want):
    ok = got == want or (isinstance(got, float) and isinstance(want, float)
                         and abs(got - want) < 1e-9)
    print(f"  [{'ok  ' if ok else 'FAIL'}] {label}: got={got!r} want={want!r}")
    return ok


def _run(rid, gold, a_pages, b_pages, qtype="factoid", ctype="prose"):
    """Synthetic A/B log rows. gold/pages are {doc: [pages]} maps."""
    common = {"question_type": qtype, "content_type": ctype, "lang": "en",
              "question_id": rid, "question": f"Q for {rid}?", "gold_pages": gold}
    return [
        {"run_id": rid, "phase": "vector", **common,
         "evidence": {"retrieved": a_pages},
         "retrieved": [{"doc": d, "page": p} for d, ps in a_pages.items() for p in ps],
         "usage": {"in": 10, "out": 5}, "calls": 1, "latency_ms": 100},
        {"run_id": rid, "phase": "reasoning", **common,
         "evidence": {"read": b_pages, "read_order": [[d, p] for d, ps in b_pages.items() for p in ps]},
         "usage": {"in": 20, "out": 8}, "tool_calls": 3, "latency_ms": 200},
    ]


def main():
    ok = True

    # ---- unit: F1 helper
    ok &= check("f1 perfect", metrics._f1({("d", 1)}, {("d", 1)}), 1.0)
    ok &= check("f1 no overlap", metrics._f1({("d", 2)}, {("d", 1)}), 0.0)
    ok &= check("f1 none when no gold", metrics._f1({("d", 1)}, set()), None)
    # gold={1,2}, hit={1,2,3}: P=2/3, R=1, F1=2*(2/3)/(2/3+1)=0.8
    ok &= check("f1 mixed", round(metrics._f1({("d", 1), ("d", 2), ("d", 3)},
                                              {("d", 1), ("d", 2)}), 4), 0.8)

    # ---- per_run wiring: A retrieves gold p.5 plus a miss; B reads exactly gold
    rows = _run("r1", {"D.pdf": [5]}, {"D.pdf": [5, 9]}, {"D.pdf": [5]})
    rec = metrics.per_run(rows)[0]
    ok &= check("A recall", rec["A"]["recall"], 1.0)
    ok &= check("A precision (1 of 2 pages gold)", rec["A"]["precision"], 0.5)
    ok &= check("A f1", round(rec["A"]["f1"], 4), round(2 * 1.0 * 0.5 / 1.5, 4))
    ok &= check("B recall/precision/f1 perfect",
                (rec["B"]["recall"], rec["B"]["precision"], rec["B"]["f1"]), (1.0, 1.0, 1.0))

    # ---- unanswerable control: gold empty -> retrieval metrics are None
    rows2 = _run("r2", {}, {"D.pdf": [3]}, {"D.pdf": [4]})
    rec2 = metrics.per_run(rows2)[0]
    ok &= check("no-gold run flagged", rec2["has_gold"], False)
    ok &= check("A f1 None when no gold", rec2["A"]["f1"], None)

    # ---- aggregation: control excluded from the f1 mean, A f1 present
    summ = metrics.summarize(rows + rows2)
    ok &= check("two runs aggregated", summ["n_runs"], 2)
    ok &= check("one unanswerable", summ["n_unanswerable"], 1)
    ok &= check("A f1 mean over the 1 answerable run",
                round(summ["overall"]["A"]["f1"]["mean"], 4), round(2 * 1.0 * 0.5 / 1.5, 4))
    ok &= check("A f1 n excludes the control", summ["overall"]["A"]["f1"]["n"], 1)

    # ---- per_run carries the drill-down fields (question, content_type, lang)
    rec = metrics.per_run(rows)[0]
    ok &= check("per_run has question text", rec["question"], "Q for r1?")
    ok &= check("per_run has content_type", rec["content_type"], "prose")
    ok &= check("per_run has lang", rec["lang"], "en")

    # ---- CSV mirrors the popup: popup column order, one row per question,
    #      plus per-type and overall totals rows
    import csv as _csv
    csv_text = metrics.to_csv(rows + rows2)
    parsed = list(_csv.DictReader(io.StringIO(csv_text)))
    header = csv_text.splitlines()[0].split(",")
    ok &= check("CSV columns match the popup order",
                header, ["question_id", "question", "question_type", "content_type",
                         "gold_pages", "lang", "A_recall", "B_recall", "A_precision",
                         "B_precision", "A_f1", "B_f1", "A_mrr", "B_reads_to_gold", "judge"])
    q_rows = [r for r in parsed if r["question_id"] != "__TOTAL__"]
    tot_rows = [r for r in parsed if r["question_id"] == "__TOTAL__"]
    ok &= check("one row per question (r1 answerable + r2 control)", len(q_rows), 2)
    ok &= check("totals rows present (per-type + overall)", len(tot_rows) >= 2, True)
    r1row = next(r for r in q_rows if r["question_id"] == "r1")
    ok &= check("per-question A_f1 written", round(float(r1row["A_f1"]), 4), round(2 * 1.0 * 0.5 / 1.5, 4))
    overall = next(r for r in tot_rows if r["question_type"] == "ALL")
    ok &= check("overall totals A_f1 = mean over answerable runs",
                round(float(overall["A_f1"]), 4), round(2 * 1.0 * 0.5 / 1.5, 4))

    # ---- Excel export: three sheets, answers carried through
    rows_with_ans = [dict(r, answer=("A-ans" if r["phase"] == "vector" else "B-ans"))
                     for r in (rows + rows2)]
    xlsx = metrics.to_xlsx(rows_with_ans)
    ok &= check("xlsx is a non-trivial zip (PK header)", xlsx[:2], b"PK")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx))
    ok &= check("three sheets in order", wb.sheetnames, ["Summary", "Detail", "Answers"])
    ok &= check("Summary header starts as dashboard",
                [c.value for c in wb["Summary"][1]][:3], ["Type", "n", "Recall@k A"])
    ans = wb["Answers"]
    hdr = [c.value for c in ans[1]]
    first = {hdr[i]: c.value for i, c in enumerate(ans[2])}
    ok &= check("Answers sheet carries both pipelines' text",
                (first["A_answer"], first["B_answer"]), ("A-ans", "B-ans"))

    # ---- judge verdicts are logged as A/B/tie (judge_pair maps 1/2 back)
    jrecs = [{"judge": "A"}, {"judge": "B"}, {"judge": "B"}, {"judge": "tie"}]
    ok &= check("judge counts read A/B winners", metrics._judge_str(jrecs, as_pct=False),
                "A1·B2·t1")
    judged = rows_with_ans + [{"run_id": rows_with_ans[0]["run_id"], "phase": "judge",
                               "winner": "B"}]
    ok &= check("Answers sheet carries the judge winner",
                metrics._answer_records(judged)[0]["judge"], "B")

    print("\nRESULT:", "ALL PASS" if ok else "FAILURES PRESENT")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())

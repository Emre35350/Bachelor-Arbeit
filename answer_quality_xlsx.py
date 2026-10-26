"""Answer-quality tables for the final run, from the judge rows in the log.

Reads the `phase: "judge"` rows of the final batch (written by
`python -m app.judge_batch --round N`) and writes one workbook:

  Run 1 … Run N    the four tables for that judge round, an Excel-formula check,
                   and every verdict of the round (scores, rationale, both answers)
  Mean of N runs   the per-run headline numbers with their mean ± sd, and the four
                   tables recomputed on the combined verdicts (see `combine`)
  Agreement        how consistently the judge rounds agree with each other

The four tables:
  1) pairwise preference (A win / B win / tie), per test case and per question
     (majority over the three repeats), with a two-sided sign test on questions;
  2) the judge's 1-5 rubric scores (faithfulness, comprehensiveness, correctness) per system,
     with a Wilcoxon signed-rank test on the per-question mean difference;
  3) / 4) the same broken down by question type and by content type.

    python answer_quality_xlsx.py      -> ..\\answer-quality.xlsx (+ tables on stdout)
"""
import math
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy.stats import binomtest, wilcoxon

from app import logstore
from app.judge_batch import judge_round

FINAL = "batch-46ded84d"
OUT = Path(__file__).resolve().parent.parent / "answer-quality.xlsx"
CRITERIA = ("grounded", "complete", "correct")   # keys in the judge prompt and the log
# reported under the names these criteria have in the RAG-evaluation literature:
# faithfulness (Es et al., RAGAS; Saad-Falcon et al., ARES), comprehensiveness
# (Edge et al., GraphRAG) and answer correctness; the definitions are unchanged
LABEL = {"grounded": "Faithfulness", "complete": "Comprehensiveness", "correct": "Correctness"}
WINNERS = ("A", "B", "tie")

HEAD_FILL = PatternFill("solid", fgColor="3B5B33")
TOT_FILL = PatternFill("solid", fgColor="F1F5F2")
WIN_FILL = PatternFill("solid", fgColor="C6EFCE")      # light green: the better system
THIN = Side(style="thin", color="D0D5CC")
NOTE = Font(size=9, italic=True, color="5B646E")


# ------------------------------------------------------------------ records
def judge_rounds(rows: list) -> dict[int, list[dict]]:
    """{round: [verdict record]} for the final batch, latest verdict per run_id and
    round, each enriched with the run's metadata and both answers."""
    meta = {r["run_id"]: r for r in rows
            if r.get("batch_id") == FINAL and r.get("phase") == "vector"}
    b_answer = {r["run_id"]: r.get("answer", "") for r in rows
                if r.get("batch_id") == FINAL and r.get("phase") == "reasoning"}
    out: dict[int, dict] = defaultdict(dict)
    for r in rows:
        if r.get("phase") == "judge" and r.get("run_id") in meta:
            m = meta[r["run_id"]]
            # unanswerable controls are labelled "prose" in the question file;
            # report them separately so prose/table cover the 59 + 22 answerable ones
            ctype = m.get("content_type") if m.get("gold_pages") else "unanswerable"
            out[judge_round(r)][r["run_id"]] = {
                **r, "question_type": m.get("question_type"), "content_type": ctype,
                "question_id": m.get("question_id"), "repeat_index": m.get("repeat_index"),
                "question": m.get("question"), "A_answer": m.get("answer", ""),
                "B_answer": b_answer.get(r["run_id"], "")}
    return {k: sorted(v.values(), key=lambda r: r["run_id"]) for k, v in sorted(out.items())}


def combine(rounds: dict[int, list[dict]]) -> list[dict]:
    """One record per test case judged in every round: rubric scores are the mean
    over the rounds, and the verdict is split into weights (e.g. B won 2 of 3
    rounds -> B 0.67, A 0.33), so win shares equal the mean of the rounds' shares."""
    by_run = [{r["run_id"]: r for r in recs} for recs in rounds.values()]
    common = sorted(set.intersection(*(set(d) for d in by_run)))
    out = []
    for rid in common:
        group = [d[rid] for d in by_run]
        w = {k: sum(1 for g in group if g["winner"] == k) / len(group) for k in WINNERS}
        scores = {who: {c: sum(score(g, who, c) for g in group) / len(group) for c in CRITERIA}
                  for who in ("A", "B")
                  if all(score(g, who, c) is not None for g in group for c in CRITERIA)}
        out.append({**group[0], "w": w, "scores": scores, "winner": None})
    return out


def weights(rec: dict) -> dict:
    return rec.get("w") or {rec["winner"]: 1.0}


def score(rec: dict, who: str, crit: str | None = None):
    s = (rec.get("scores") or {}).get(who) or {}
    if crit:
        v = s.get(crit)
        return float(v) if isinstance(v, (int, float)) else None
    vals = [s.get(c) for c in CRITERIA if isinstance(s.get(c), (int, float))]
    return sum(vals) / len(vals) if len(vals) == len(CRITERIA) else None


def mean_sd(xs):
    xs = [x for x in xs if x is not None]
    if not xs:
        return None, None
    m = sum(xs) / len(xs)
    return m, (math.sqrt(sum((x - m) ** 2 for x in xs) / len(xs)) if len(xs) > 1 else 0.0)


def shares(recs):
    n = len(recs)
    c = {k: sum(weights(r).get(k, 0.0) for r in recs) for k in WINNERS}
    return n, {k: (v / n if n else None) for k, v in c.items()}, c


def per_question(recs):
    """Majority verdict and mean score difference (B - A) per question."""
    by_q = defaultdict(list)
    for r in recs:
        by_q[r["question_id"]].append(r)
    maj, diffs = [], []
    for sub in by_q.values():
        a = sum(weights(r).get("A", 0.0) for r in sub)
        b = sum(weights(r).get("B", 0.0) for r in sub)
        maj.append("A" if a > b else "B" if b > a else "tie")
        d = [score(r, "B") - score(r, "A") for r in sub
             if score(r, "A") is not None and score(r, "B") is not None]
        if d:
            diffs.append(sum(d) / len(d))
    return maj, diffs


def first_shown_share(recs):
    """(wins of the answer shown first, decisive verdicts) — the position-bias check."""
    decisive = [r for r in recs if r.get("winner") in ("A", "B")]
    first = sum(1 for r in decisive if (r["winner"] == "A") == (r.get("a_was_answer") == 1))
    return first, len(decisive)


def headline(recs) -> dict:
    """The numbers compared across rounds (shares as 0-1, scores on the 1-5 scale)."""
    n, sh, _ = shares(recs)
    maj, _ = per_question(recs)
    h = {"A wins (share of test cases)": sh["A"], "B wins (share of test cases)": sh["B"],
         "Ties (share of test cases)": sh["tie"],
         "Questions won by A": sum(1 for m in maj if m == "A"),
         "Questions won by B": sum(1 for m in maj if m == "B")}
    for c in CRITERIA + (None,):
        label = LABEL[c] if c else "Mean of the three"
        h[f"{label} — A"] = mean_sd([score(r, "A", c) for r in recs])[0]
        h[f"{label} — B"] = mean_sd([score(r, "B", c) for r in recs])[0]
    first, dec = first_shown_share(recs)
    h["Decisive verdicts for the answer shown first"] = first / dec if dec else None
    return h


# ------------------------------------------------------------------ formatting
def pct(v):
    # one decimal, so the shares of a row add up to 100% instead of 99%/101%
    return "—" if v is None else f"{100 * v:.1f}%"


def num(v, sd=None, dp=2):
    if v is None:
        return "—"
    return f"{v:.{dp}f}" + (f" ± {sd:.{dp}f}" if sd is not None else "")


def cnt(v):
    return str(int(v)) if float(v).is_integer() else f"{v:.1f}"


def md_table(head, body):
    lines = ["| " + " | ".join(head) + " |", "|" + "|".join("---" for _ in head) + "|"]
    lines += ["| " + " | ".join(str(c) for c in row) + " |" for row in body]
    return "\n".join(lines)


# ------------------------------------------------------------------ tables
def build(recs):
    """The four tables as (title, header, rows) — shared by the workbook and stdout."""
    n, sh, c = shares(recs)
    maj, diffs = per_question(recs)
    qa, qb, qt = (sum(1 for m in maj if m == k) for k in WINNERS)
    sign_p = binomtest(qb, qa + qb, 0.5).pvalue if qa + qb else None
    wil_p = wilcoxon(diffs).pvalue if any(d != 0 for d in diffs) else None

    t1 = ("Table 1 — Pairwise preference",
          ["Level", "n", "A wins", "B wins", "Ties", "Sign test p (A vs B)"],
          [["Test cases (runs)", n, f"{cnt(c['A'])} ({pct(sh['A'])})",
            f"{cnt(c['B'])} ({pct(sh['B'])})", f"{cnt(c['tie'])} ({pct(sh['tie'])})", "—"],
           ["Questions (majority over the 3 repeats)", len(maj), f"{qa} ({pct(qa / len(maj))})",
            f"{qb} ({pct(qb / len(maj))})", f"{qt} ({pct(qt / len(maj))})",
            "—" if sign_p is None else f"{sign_p:.2g}"]])

    rows2 = []
    for crit in CRITERIA + (None,):
        ma, sa = mean_sd([score(r, "A", crit) for r in recs])
        mb, sb = mean_sd([score(r, "B", crit) for r in recs])
        rows2.append([LABEL[crit] if crit else "Mean of the three",
                      num(ma, sa), num(mb, sb), num(None if ma is None else mb - ma), ""])
    rows2[-1][-1] = "—" if wil_p is None else f"{wil_p:.2g}"
    t2 = ("Table 2 — Rubric scores (1–5, mean ± sd over test cases)",
          ["Criterion", "A · Vector-Based RAG", "B · Reasoning-Based Retrieval", "Δ (B − A)",
           "Wilcoxon p (per question)"], rows2)

    def breakdown(key, order=None):
        groups = defaultdict(list)
        for r in recs:
            groups[r.get(key) or "—"].append(r)
        keys = order or sorted(groups, key=lambda k: -len(groups[k]))
        body = []
        for k in keys:
            g = groups.get(k, [])
            if not g:
                continue
            gn, gs, _ = shares(g)
            body.append([k, gn, pct(gs["A"]), pct(gs["B"]), pct(gs["tie"]),
                         num(mean_sd([score(r, "A") for r in g])[0]),
                         num(mean_sd([score(r, "B") for r in g])[0]),
                         num(mean_sd([score(r, "A", "correct") for r in g])[0]),
                         num(mean_sd([score(r, "B", "correct") for r in g])[0])])
        return body

    bhead = ["n", "A wins", "B wins", "Ties", "Score A", "Score B", "Correctness A", "Correctness B"]
    t3 = ("Table 3 — By question type", ["Question type"] + bhead, breakdown("question_type"))
    t4 = ("Table 4 — By content type", ["Content type"] + bhead,
          breakdown("content_type", ["prose", "table", "unanswerable"]))
    return [t1, t2, t3, t4]


def audit_table(recs):
    first, dec = first_shown_share(recs)
    pos_p = binomtest(first, dec, 0.5).pvalue if dec else None
    tok_in = sum((r.get("usage") or {}).get("in", 0) for r in recs)
    tok_out = sum((r.get("usage") or {}).get("out", 0) for r in recs)
    return ("Judge audit", ["Check", "Value"],
            [["Judge model", f"{recs[0].get('judge_provider')}/{recs[0].get('judge_model')}"],
             ["Verdicts", f"{len(recs)} of 267 test cases"],
             ["Decisive verdicts for the answer shown first",
              f"{first}/{dec} ({pct(first / dec if dec else None)}), binomial p = {pos_p:.2g}"
              if pos_p is not None else "—"],
             ["Judge tokens (in / out)", f"{tok_in:,} / {tok_out:,}"]])


def per_run_table(rounds):
    """Headline numbers of every round side by side, with their mean ± sd."""
    hs = [headline(recs) for recs in rounds.values()]
    body = []
    for key in hs[0]:
        vals = [h[key] for h in hs]
        m, sd = mean_sd(vals)
        is_share = "share" in key or key.startswith("Decisive")
        is_count = key.startswith("Questions")
        fmt = pct if is_share else (lambda v: f"{v:.0f}") if is_count else num
        spread = ("—" if sd is None else f"± {100 * sd:.1f} pp" if is_share
                  else f"± {sd:.1f}" if is_count else f"± {sd:.2f}")
        body.append([key] + [fmt(v) for v in vals]
                    + [f"{m:.1f}" if is_count else fmt(m), spread])
    head = ["Measure"] + [f"Run {k}" for k in rounds] + ["Mean", "SD between runs"]
    return ("Headline numbers per judge run", head, body)


def agreement_tables(rounds):
    by_run = [{r["run_id"]: r for r in recs} for recs in rounds.values()]
    common = sorted(set.intersection(*(set(d) for d in by_run)))
    k = len(by_run)
    verdicts = [[d[rid]["winner"] for d in by_run] for rid in common]
    N = len(verdicts)

    same = sum(1 for v in verdicts if len(set(v)) == 1)
    split = sum(1 for v in verdicts if len(set(v)) == 2)
    t_w = [["All runs give the same winner", f"{same} ({pct(same / N)})"],
           ["Runs disagree (one run differs)" if k == 3 else "Runs disagree",
            f"{split} ({pct(split / N)})"]]
    if k >= 3:
        diff = sum(1 for v in verdicts if len(set(v)) == 3)
        t_w.append(["All three runs differ", f"{diff} ({pct(diff / N)})"])
    labels = list(rounds)
    for i in range(k):
        for j in range(i + 1, k):
            agree = sum(1 for v in verdicts if v[i] == v[j])
            t_w.append([f"Run {labels[i]} vs Run {labels[j]}: same winner",
                        f"{agree} ({pct(agree / N)})"])
    # Fleiss' kappa: chance-corrected agreement of k raters over N items
    nij = [[v.count(c) for c in WINNERS] for v in verdicts]
    p_i = [(sum(x * x for x in row) - k) / (k * (k - 1)) for row in nij]
    p_j = [sum(row[c] for row in nij) / (N * k) for c in range(len(WINNERS))]
    pe = sum(p * p for p in p_j)
    kappa = (sum(p_i) / N - pe) / (1 - pe) if pe < 1 else None
    t_w.append(["Fleiss' kappa (winner)", "—" if kappa is None else f"{kappa:.2f}"])
    t1 = ("Winner agreement between judge runs", ["Measure", "Value"], t_w)

    rows = []
    for who in ("A", "B"):
        for c in CRITERIA:
            vals = [[score(d[rid], who, c) for d in by_run] for rid in common]
            vals = [v for v in vals if None not in v]
            ident = sum(1 for v in vals if len(set(v)) == 1)
            within_sd = mean_sd([mean_sd(v)[1] for v in vals])[0]
            rng = mean_sd([max(v) - min(v) for v in vals])[0]
            rows.append([f"{who} · {LABEL[c]}", f"{ident}/{len(vals)} ({pct(ident / len(vals))})",
                         num(within_sd), num(rng)])
    t2 = ("Score stability between judge runs (per answer)",
          ["System · criterion", "Identical score in all runs", "Mean SD between runs",
           "Mean range (max − min)"], rows)
    return [t1, t2]


# ------------------------------------------------------------------ workbook
SYSTEM = {"A": "Vector-Based RAG (A)", "B": "Reasoning-Based Retrieval (B)", "tie": "Tie"}


def lead_num(v):
    """Leading number of a formatted cell: '57.7 (22%)' -> 57.7, '2.73 ± 0.82' -> 2.73."""
    m = re.match(r"\s*(-?\d+(?:\.\d+)?)", str(v))
    return float(m.group(1)) if m else None


def better(a, b):
    """'A' / 'B' for the larger of two formatted values, 'tie' if equal, None if not numeric."""
    x, y = lead_num(a), lead_num(b)
    if x is None or y is None:
        return None
    return "A" if x > y else "B" if y > x else "tie"


def write_table(ws, r, title, head, body, bold_last=False, pairs=None, winner_head=None):
    """Write one titled, bordered table starting at row r; return the next free row.

    pairs: (A column, B column) index pairs (1-based); the better value of each pair
    is filled light green. winner_head: if given, a last column names the winner of
    the first pair in every row."""
    if winner_head:
        head = head + [winner_head]
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=11)
    r += 1
    for c, h in enumerate(head, start=1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top = r
    for i, row in enumerate(body):
        r += 1
        if winner_head:
            w = better(row[pairs[0][0] - 1], row[pairs[0][1] - 1])
            row = row + [SYSTEM.get(w, "—")]
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center")
            if bold_last and i == len(body) - 1:
                cell.font, cell.fill = Font(bold=True), TOT_FILL
        for ca, cb in pairs or ():
            w = better(row[ca - 1], row[cb - 1])
            if w in ("A", "B"):
                ws.cell(row=r, column=ca if w == "A" else cb).fill = WIN_FILL
        if winner_head:
            wc = ws.cell(row=r, column=len(head))
            wc.font = Font(bold=True, color="006100" if row[-1] != "—" else "1F2328")
    for row in ws.iter_rows(min_row=top, max_row=r, max_col=len(head)):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    return r + 3


def sheet_header(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=9, color="5B646E")
    ws["A3"] = ("Criteria (literature names; rubric labels given to the judge in brackets): "
                "Faithfulness [grounded] · Comprehensiveness [complete] · Correctness [correct]")
    ws["A3"].font = Font(size=9, italic=True, color="5B646E")
    ws.column_dimensions["A"].width = 40
    for c in range(2, 18):
        ws.column_dimensions[get_column_letter(c)].width = 16


RAW_COLS = ["run_id", "question_id", "question_type", "content_type", "repeat", "winner",
            *(f"{w} {LABEL[c].lower()}" for w in "AB" for c in CRITERIA),
            "shown first", "rationale", "question", "A answer", "B answer"]


def write_run_sheet(wb, rnd, recs):
    ws = wb.create_sheet(f"Run {rnd}")
    sheet_header(ws, f"Answer quality — judge run {rnd}",
                 f"batch id: {FINAL} · judged post hoc on the logged answers · "
                 "labels hidden, answer order randomized per verdict")
    r = 4
    for t in build(recs) + [audit_table(recs)]:
        r = write_table(ws, r, *t, bold_last=t[0].startswith("Table 2"))

    # raw verdicts first, so the formula block can point at their rows
    check_rows = 12
    raw_title = r + check_rows + 4
    first, last = raw_title + 2, raw_title + 1 + len(recs)
    rng = lambda col: f"{col}{first}:{col}{last}"

    ws.cell(row=r, column=1, value="Excel check — the same numbers recomputed by Excel "
            "formulas over the verdicts listed below").font = Font(bold=True, size=11)
    head = ["Check", "Python (tables above)", "Excel formula"]
    checks = [("A wins", f'=COUNTIF({rng("F")},"A")', shares(recs)[2]["A"]),
              ("B wins", f'=COUNTIF({rng("F")},"B")', shares(recs)[2]["B"]),
              ("Ties", f'=COUNTIF({rng("F")},"tie")', shares(recs)[2]["tie"])]
    for col, (who, c) in zip("GHIJKL", [(w, c) for w in "AB" for c in CRITERIA]):
        checks.append((f"Mean {LABEL[c].lower()} — {who}", f"=AVERAGE({rng(col)})",
                       mean_sd([score(x, who, c) for x in recs])[0]))
    checks += [("Mean of the three — A", f"=AVERAGE(G{first}:I{last})",
                mean_sd([score(x, "A") for x in recs])[0]),
               ("Mean of the three — B", f"=AVERAGE(J{first}:L{last})",
                mean_sd([score(x, "B") for x in recs])[0])]
    rr = r + 1
    for c, h in enumerate(head, start=1):
        cell = ws.cell(row=rr, column=c, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF", size=10), HEAD_FILL
    for label, formula, py in checks:
        rr += 1
        ws.cell(row=rr, column=1, value=label)
        ws.cell(row=rr, column=2, value=round(py, 4)).number_format = "0.00"
        ws.cell(row=rr, column=3, value=formula).number_format = "0.00"
    for row in ws.iter_rows(min_row=r + 1, max_row=rr, max_col=3):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if cell.column > 1:
                cell.alignment = Alignment(horizontal="center")
    assert rr < raw_title, "formula block overlaps the verdict list"

    ws.cell(row=raw_title, column=1,
            value=f"All {len(recs)} verdicts of run {rnd}").font = Font(bold=True, size=11)
    for c, h in enumerate(RAW_COLS, start=1):
        cell = ws.cell(row=raw_title + 1, column=c, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF", size=10), HEAD_FILL
    for i, x in enumerate(recs):
        sa, sb = (x.get("scores") or {}).get("A", {}), (x.get("scores") or {}).get("B", {})
        shown = "A" if x.get("a_was_answer") == 1 else "B"
        vals = [x["run_id"], x["question_id"], x["question_type"], x["content_type"],
                x.get("repeat_index"), x["winner"],
                *(sa.get(c) for c in CRITERIA), *(sb.get(c) for c in CRITERIA),
                shown, x.get("rationale", ""), x.get("question", ""),
                x.get("A_answer", ""), x.get("B_answer", "")]
        for c, v in enumerate(vals, start=1):
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)   # control chars break the xlsx
            ws.cell(row=first + i, column=c, value=v)
    ws.auto_filter.ref = f"A{raw_title + 1}:{get_column_letter(len(RAW_COLS))}{last}"
    for col, w in zip("NOPQ", (60, 50, 60, 60)):
        ws.column_dimensions[col].width = w


NAMED_HEAD = {"A wins": "A wins · Vector-Based RAG", "B wins": "B wins · Reasoning-Based Retrieval",
              "Score A": "Score A · Vector", "Score B": "Score B · Reasoning",
              "Correctness A": "Correctness A · Vector", "Correctness B": "Correctness B · Reasoning"}


def verdict_banner(ws, r, tables, ncols=10):
    """Plain-language result box: who won, and on which evidence."""
    t1, t2, t3, t4 = tables
    tc, q = t1[2]
    winner = better(tc[2], tc[3])
    other = "A" if winner == "B" else "B"
    crit = [better(row[1], row[2]) for row in t2[2][:3]]
    n_crit = sum(1 for w in crit if w == winner)
    n_type = sum(1 for row in t3[2] if better(row[2], row[3]) == winner)
    n_ct = sum(1 for row in t4[2] if better(row[2], row[3]) == winner)
    win_idx = 2 if winner == "A" else 3       # t1 columns: level, n, A wins, B wins, ties, p
    other_col = 2 if other == "A" else 3      # win-share column of the runner-up in t3/t4
    closest = max(t4[2], key=lambda row: lead_num(row[other_col]) or 0)
    lines = [
        (f"Result: {SYSTEM[winner]} gives the better answers", Font(bold=True, size=14, color="006100")),
        (f"• Preferred by the judge in {tc[win_idx].split('(')[1].rstrip(')')} of the 267 test cases "
         f"and in {q[win_idx].split(' (')[0]} of the 89 questions (sign test p = {q[5]}).", Font(size=11)),
        (f"• Higher rubric score on {n_crit} of 3 criteria — overall "
         f"{t2[2][3][1 if winner == 'A' else 2].split(' ±')[0]} vs "
         f"{t2[2][3][2 if winner == 'A' else 1].split(' ±')[0]} on the 1–5 scale "
         f"(Wilcoxon p = {t2[2][3][4]}).", Font(size=11)),
        (f"• Wins in {n_type} of {len(t3[2])} question types and {n_ct} of {len(t4[2])} content "
         f"types. {SYSTEM[other]} comes closest on {closest[0]} questions "
         f"({closest[other_col]} wins).", Font(size=11)),
        ("Green cell = the better of the two systems in that row (or column). The 'Winner' column "
         "names the system with the higher value of the first compared pair.", NOTE),
    ]
    top = r
    for text, font in lines:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 24 if r == top else 17
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="EAF6EC")
        r += 1
    for row in ws.iter_rows(min_row=top, max_row=r - 1, max_col=ncols):
        for cell in row:
            cell.border = Border(top=THIN if cell.row == top else None,
                                 bottom=THIN if cell.row == r - 1 else None,
                                 left=THIN if cell.column == 1 else None,
                                 right=THIN if cell.column == ncols else None)
    return r + 2


def mark_row_pairs(ws, first_row, body, ncols_values, winner_col):
    """Per-run table: rows come in A/B pairs ('… — A' / '… — B', 'A wins' / 'B wins');
    fill the better value of each pair green per column and name the winner (by mean)."""
    rows = {label: first_row + i for i, (label, *_) in enumerate(body)}
    pair_keys = [("A wins (share of test cases)", "B wins (share of test cases)"),
                 ("Questions won by A", "Questions won by B")]
    pair_keys += [(l, l[:-1] + "B") for l in rows if l.endswith(" — A")]
    for ka, kb in pair_keys:
        ra, rb = rows[ka], rows[kb]
        for c in range(2, 2 + ncols_values):
            w = better(ws.cell(row=ra, column=c).value, ws.cell(row=rb, column=c).value)
            if w in ("A", "B"):
                ws.cell(row=ra if w == "A" else rb, column=c).fill = WIN_FILL
        mean_col = 1 + ncols_values
        w = better(ws.cell(row=ra, column=mean_col).value, ws.cell(row=rb, column=mean_col).value)
        ws.merge_cells(start_row=ra, start_column=winner_col, end_row=rb, end_column=winner_col)
        cell = ws.cell(row=ra, column=winner_col, value=SYSTEM.get(w, "—"))
        cell.font = Font(bold=True, color="006100")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_mean_sheet(wb, rounds, combined):
    k = len(rounds)
    ws = wb.create_sheet(f"Mean of {k} runs")
    sheet_header(ws, f"Answer quality — mean of {k} judge runs",
                 f"batch id: {FINAL} · {len(combined)} test cases judged in every run · "
                 "A = Vector-Based RAG, B = Reasoning-Based Retrieval")
    tables = build(combined)
    r = verdict_banner(ws, 4, tables)

    title, head, body = per_run_table(rounds)
    start = r
    r = write_table(ws, r, title, head + ["Better system (mean)"], [row + [""] for row in body])
    mark_row_pairs(ws, start + 2, body, ncols_values=k + 1, winner_col=len(head) + 1)
    ws.cell(row=r - 1, column=1, value=(
        "The tables below use the combined verdicts: each test case's rubric scores are the "
        f"mean over the {k} runs, and its verdict counts as a fraction of a win per run (e.g. "
        f"B won 2 of {k} runs → 0.67 B), so the win shares equal the mean of the runs' shares. "
        "Question-level majority, sign test and Wilcoxon test are computed on these combined "
        "values.")).font = NOTE
    r += 1
    t1, t2, t3, t4 = tables
    named = lambda h: [NAMED_HEAD.get(x, x) for x in h]
    r = write_table(ws, r, t1[0], named(t1[1]), t1[2], pairs=[(3, 4)], winner_head="Winner")
    r = write_table(ws, r, t2[0], t2[1], t2[2], bold_last=True, pairs=[(2, 3)],
                    winner_head="Higher score")
    for t in (t3, t4):
        r = write_table(ws, r, t[0], named(t[1]), t[2], pairs=[(3, 4), (6, 7), (8, 9)],
                        winner_head="Winner (by win share)")
    for col in "FGJ":                          # the 'Winner' columns
        ws.column_dimensions[col].width = 28


def write_agreement_sheet(wb, rounds):
    ws = wb.create_sheet("Agreement")
    sheet_header(ws, "Answer quality — agreement between judge runs",
                 f"batch id: {FINAL} · same answer pairs, answer order re-randomized per run")
    r = 4
    for t in agreement_tables(rounds):
        r = write_table(ws, r, *t)
    ws.cell(row=r - 1, column=1, value=(
        "Fleiss' kappa corrects the raw agreement for the agreement expected by chance from how "
        "often each verdict occurs; 0.61–0.80 is conventionally read as substantial and above "
        "0.80 as almost perfect agreement (Landis & Koch, 1977). Because B wins most verdicts, "
        "chance agreement is high and kappa is lower than the raw agreement.")).font = NOTE


def main():
    rounds = judge_rounds(logstore.read_all())
    if not rounds:
        raise SystemExit(f"No judge rows for {FINAL}. Run: python -m app.judge_batch --batch {FINAL}")
    wb = Workbook()
    wb.remove(wb.active)
    for rnd, recs in rounds.items():
        write_run_sheet(wb, rnd, recs)
        print(f"run {rnd}: {len(recs)} verdicts")
    if len(rounds) > 1:
        combined = combine(rounds)
        write_mean_sheet(wb, rounds, combined)
        write_agreement_sheet(wb, rounds)
        for title, head, body in [per_run_table(rounds)] + build(combined) + agreement_tables(rounds):
            print(f"\n### {title}\n")
            print(md_table(head, body))
    else:
        for title, head, body in build(rounds[1]):
            print(f"\n### {title}\n")
            print(md_table(head, body))
    wb.save(OUT)
    print("\nwritten:", OUT)


if __name__ == "__main__":
    main()

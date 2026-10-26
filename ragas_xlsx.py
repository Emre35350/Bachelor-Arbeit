"""Answer-quality tables with the RAGAS metrics (Es et al., 2024) for the final run.

Reads the `phase: "ragas"` rows written by `python -m app.ragas_eval --round N`
and writes "..\\Answer Quality with New Metrics.xlsx":

  Run 1 … Run N     the four tables of that run, an Excel-formula check and every
                    score of the run
  Mean of N Runs    the per-run headline numbers with their mean ± sd, the four
                    tables recomputed on the per-answer mean over the runs, and the
                    run-to-run stability of the scores — the better system green

The four tables (Faithfulness F, Answer relevance AR, Context relevance CR, all 0-1):
  1) overall scores on the answerable questions, with a Wilcoxon signed-rank test
     on the per-question mean difference;
  2) per metric, on how many questions each system scored higher, with a sign test;
  3) / 4) the scores by question type and by content type.
Unanswerable controls are excluded from tables 1 and 2: RAGAS defines no rule for
an answer that correctly abstains, so its scores there do not measure quality.

    python ragas_xlsx.py
"""
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.stats import binomtest, wilcoxon

from answer_quality_xlsx import (HEAD_FILL, NOTE, THIN, WIN_FILL, better, md_table, mean_sd,
                                 write_table)
from app import logstore
from app.ragas_eval import EMBED_MODEL, EVAL_MODEL, ragas_round

FINAL = "batch-46ded84d"
ROUNDS = (1, 2, 3)
OUT = Path(__file__).resolve().parent.parent / "Answer Quality with New Metrics.xlsx"
METRICS = ("faithfulness", "answer_relevance", "context_relevance")
NAME = {"faithfulness": "Faithfulness", "answer_relevance": "Answer relevance",
        "context_relevance": "Context relevance"}
SHORT = {"faithfulness": "F", "answer_relevance": "AR", "context_relevance": "CR"}
SYSTEM = {"A": "Vector-Based RAG (A)", "B": "Reasoning-Based Retrieval (B)", "tie": "Tie"}


# ------------------------------------------------------------------ records
def load_rounds(rows: list) -> dict[int, list[dict]]:
    """{round: [test case]}; a test case holds both systems' scores and details."""
    out: dict[int, dict] = defaultdict(dict)
    for r in rows:
        if r.get("phase") != "ragas" or r.get("batch_id") != FINAL or ragas_round(r) not in ROUNDS:
            continue
        tc = out[ragas_round(r)].setdefault(r["run_id"], {
            "run_id": r["run_id"], "question_id": r["question_id"],
            "question_type": r["question_type"], "repeat": r["repeat_index"],
            "answerable": bool(r.get("gold_pages")), "question": r.get("question"),
            "content_type": r["content_type"] if r.get("gold_pages") else "unanswerable"})
        d = r.get("detail") or {}
        tc[r["system"]] = {
            **{m: r.get(m) for m in METRICS},
            "n_statements": len(d.get("faithfulness", {}).get("statements", [])),
            "n_supported": sum(1 for v in d.get("faithfulness", {}).get("verdicts", [])
                               if str(v.get("verdict", "")).lower().startswith("yes")),
            "n_sentences": d.get("context_relevance", {}).get("n_sentences"),
            "n_selected": len(d.get("context_relevance", {}).get("selected", [])),
            "usage": r.get("usage") or {}}
    return {k: sorted((t for t in v.values() if "A" in t and "B" in t), key=lambda t: t["run_id"])
            for k, v in sorted(out.items())}


def combine(rounds: dict[int, list[dict]]) -> list[dict]:
    """Per test case and system, each metric's mean over the runs that scored it."""
    by = [{t["run_id"]: t for t in recs} for recs in rounds.values()]
    common = sorted(set.intersection(*(set(d) for d in by)))
    out = []
    for rid in common:
        g = [d[rid] for d in by]
        rec = {k: v for k, v in g[0].items() if k not in ("A", "B")}
        for s in "AB":
            rec[s] = {}
            for m in METRICS:
                vals = [x[s][m] for x in g if x[s][m] is not None]
                rec[s][m] = sum(vals) / len(vals) if vals else None
        out.append(rec)
    return out


def val(t, s, m):
    return t[s].get(m)


# ------------------------------------------------------------------ formatting
# Scores (0-1) are shown as percentages with two decimals (0.8912 -> 89.12%);
# differences between two percentages (Δ, spreads) are percentage points (pp).
def fp(v, sd=None):
    if v is None:
        return "—"
    return f"{100 * v:.2f}%" + (f" ± {100 * sd:.2f}" if sd is not None else "")


def pp(v):
    return "—" if v is None else f"{100 * v:.2f} pp"


def pct2(v):
    return "—" if v is None else f"{100 * v:.2f}%"


def pfmt(p):
    return "—" if p is None or (isinstance(p, float) and math.isnan(p)) else f"{p:.2g}"


# ------------------------------------------------------------------ statistics
def per_question_diffs(recs, m):
    """Per answerable question: mean over its repeats of (B − A) for metric m."""
    by_q = defaultdict(list)
    for t in recs:
        if t["answerable"] and val(t, "A", m) is not None and val(t, "B", m) is not None:
            by_q[t["question_id"]].append(val(t, "B", m) - val(t, "A", m))
    return {q: sum(d) / len(d) for q, d in by_q.items()}


def build(recs):
    """The four tables as (title, header, rows)."""
    ans = [t for t in recs if t["answerable"]]
    rows1, rows2 = [], []
    for m in METRICS:
        ma, sa = mean_sd([val(t, "A", m) for t in ans])
        mb, sb = mean_sd([val(t, "B", m) for t in ans])
        n = sum(1 for t in ans if val(t, "A", m) is not None and val(t, "B", m) is not None)
        diffs = list(per_question_diffs(recs, m).values())
        p = wilcoxon(diffs).pvalue if any(d != 0 for d in diffs) else None
        rows1.append([NAME[m], n, fp(ma, sa), fp(mb, sb),
                      pp(None if ma is None else mb - ma), pfmt(p)])
        qa = sum(1 for d in diffs if d < 0)
        qb = sum(1 for d in diffs if d > 0)
        qt = sum(1 for d in diffs if d == 0)
        ps = binomtest(qb, qa + qb, 0.5).pvalue if qa + qb else None
        nq = len(diffs)
        rows2.append([NAME[m], nq, f"{qa} ({pct2(qa / nq)})", f"{qb} ({pct2(qb / nq)})",
                      f"{qt} ({pct2(qt / nq)})", pfmt(ps)])
    t1 = (f"Table 1 — Overall scores, answerable questions (%, mean ± sd over test cases)",
          ["Metric", "n", "A · Vector-Based RAG", "B · Reasoning-Based Retrieval", "Δ (B − A, pp)",
           "Wilcoxon p (per question)"], rows1)
    t2 = ("Table 2 — Questions on which each system scored higher (answerable, 3 repeats averaged)",
          ["Metric", "Questions", "A higher", "B higher", "Equal", "Sign test p (A vs B)"], rows2)

    def breakdown(key, order=None):
        groups = defaultdict(list)
        for t in recs:
            groups[t[key]].append(t)
        keys = order or sorted(groups, key=lambda k: -len(groups[k]))
        body = []
        for k in keys:
            g = groups.get(k, [])
            if g:
                row = [k, len(g)]
                for m in METRICS:
                    row += [fp(mean_sd([val(t, "A", m) for t in g])[0]),
                            fp(mean_sd([val(t, "B", m) for t in g])[0])]
                body.append(row)
        return body

    head = ["n"] + [f"{SHORT[m]} {s}" for m in METRICS for s in "AB"]
    t3 = ("Table 3 — By question type (mean, %)", ["Question type"] + head,
          breakdown("question_type"))
    t4 = ("Table 4 — By content type (mean, %)", ["Content type"] + head,
          breakdown("content_type", ["prose", "table", "unanswerable"]))
    return [t1, t2, t3, t4]


def headline(recs) -> dict:
    ans = [t for t in recs if t["answerable"]]
    h = {}
    for m in METRICS:
        for s in "AB":
            h[f"{NAME[m]} — {s}"] = mean_sd([val(t, s, m) for t in ans])[0]
    return h


def per_run_table(rounds):
    hs = [headline(r) for r in rounds.values()]
    body = []
    for key in hs[0]:
        vals = [h[key] for h in hs]
        m, sd = mean_sd(vals)
        body.append([key] + [fp(v) for v in vals] + [fp(m), f"± {pp(sd)}"])
    return ("Headline numbers per run (answerable questions, mean, %)",
            ["Measure"] + [f"Run {k}" for k in rounds] + ["Mean", "SD between runs"], body)


def stability_table(rounds):
    by = [{t["run_id"]: t for t in recs} for recs in rounds.values()]
    common = sorted(set.intersection(*(set(d) for d in by)))
    rows = []
    for s in "AB":
        for m in METRICS:
            vs = [[d[rid][s][m] for d in by] for rid in common]
            vs = [v for v in vs if None not in v]
            ident = sum(1 for v in vs if max(v) - min(v) < 1e-9)
            rows.append([f"{s} · {NAME[m]}", f"{ident}/{len(vs)} ({pct2(ident / len(vs))})",
                         pp(mean_sd([mean_sd(v)[1] for v in vs])[0]),
                         pp(mean_sd([max(v) - min(v) for v in vs])[0])])
    return ("Table 5 — Run-to-run stability of the scores (per answer, all 267 test cases)",
            ["System · metric", "Identical score in all runs", "Mean SD between runs (pp)",
             "Mean range, max − min (pp)"], rows)


def audit_table(recs):
    n_ans = sum(1 for t in recs if t["answerable"])
    missing = {m: sum(1 for t in recs for s in "AB" if t[s][m] is None) for m in METRICS}
    tin = sum(t[s]["usage"].get("in", 0) for t in recs for s in "AB")
    tout = sum(t[s]["usage"].get("out", 0) for t in recs for s in "AB")
    return ("Evaluation audit", ["Check", "Value"],
            [["Evaluator LLM / embedding model", f"{EVAL_MODEL} / {EMBED_MODEL}"],
             ["Test cases scored (both systems)", f"{len(recs)} of 267 ({n_ans} answerable)"],
             ["Scores not defined (no statement / no context)",
              ", ".join(f"{SHORT[m]}: {missing[m]}" for m in METRICS)],
             ["Evaluator tokens (in / out)", f"{tin:,} / {tout:,}"]])


# ------------------------------------------------------------------ workbook
def sheet_header(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=9, color="5B646E")
    ws["A3"] = ("Metrics as defined in Es et al. (2024), Ragas: Automated Evaluation of Retrieval "
                "Augmented Generation · scores shown as % (score × 100), differences in percentage points (pp) "
                "· higher is better for all three · A = Vector-Based RAG, B = Reasoning-Based Retrieval")
    ws["A3"].font = Font(size=9, italic=True, color="5B646E")
    ws.column_dimensions["A"].width = 34
    for c in range(2, 24):
        ws.column_dimensions[get_column_letter(c)].width = 15


RAW = ["run_id", "question_id", "question_type", "content_type", "repeat", "answerable",
       "F A", "F B", "AR A", "AR B", "CR A", "CR B",
       "A statements", "A supported", "B statements", "B supported",
       "A context sentences", "A extracted", "B context sentences", "B extracted", "question"]
VAL_COLS = {("faithfulness", "A"): "G", ("faithfulness", "B"): "H",
            ("answer_relevance", "A"): "I", ("answer_relevance", "B"): "J",
            ("context_relevance", "A"): "K", ("context_relevance", "B"): "L"}


def write_run_sheet(wb, rnd, recs):
    ws = wb.create_sheet(f"Run {rnd}")
    sheet_header(ws, f"Answer quality with RAGAS metrics — run {rnd}",
                 f"batch id: {FINAL} · each answer scored post hoc against the context its "
                 "pipeline actually saw")
    r = 5
    for t in build(recs) + [audit_table(recs)]:
        r = write_table(ws, r, *t)

    n_checks = 1 + len(VAL_COLS)
    raw_title = r + n_checks + 4
    first, last = raw_title + 2, raw_title + 1 + len(recs)
    ws.cell(row=r, column=1, value="Excel check — Table 1 means recomputed by Excel formulas "
            "over the scores listed below (answerable rows)").font = Font(bold=True, size=11)
    rr = r + 1
    for c, h in enumerate(["Check", "Python (Table 1)", "Excel formula"], start=1):
        cell = ws.cell(row=rr, column=c, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF", size=10), HEAD_FILL
    ans = [t for t in recs if t["answerable"]]
    for (m, s), col in VAL_COLS.items():
        rr += 1
        ws.cell(row=rr, column=1, value=f"{NAME[m]} — {s}")
        py = mean_sd([val(t, s, m) for t in ans])[0]
        ws.cell(row=rr, column=2, value=py).number_format = "0.00%"
        ws.cell(row=rr, column=3, value=(f'=AVERAGEIFS({col}{first}:{col}{last},'
                                         f'F{first}:F{last},"yes")')).number_format = "0.00%"
    for row in ws.iter_rows(min_row=r + 1, max_row=rr, max_col=3):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if cell.column > 1:
                cell.alignment = Alignment(horizontal="center")

    ws.cell(row=raw_title, column=1, value=f"All {len(recs)} test cases of run {rnd}").font = \
        Font(bold=True, size=11)
    for c, h in enumerate(RAW, start=1):
        cell = ws.cell(row=raw_title + 1, column=c, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF", size=10), HEAD_FILL
    for i, t in enumerate(recs):
        A, B = t["A"], t["B"]
        vals = [t["run_id"], t["question_id"], t["question_type"], t["content_type"], t["repeat"],
                "yes" if t["answerable"] else "no",
                A["faithfulness"], B["faithfulness"], A["answer_relevance"], B["answer_relevance"],
                A["context_relevance"], B["context_relevance"],
                A["n_statements"], A["n_supported"], B["n_statements"], B["n_supported"],
                A["n_sentences"], A["n_selected"], B["n_sentences"], B["n_selected"],
                ILLEGAL_CHARACTERS_RE.sub("", t.get("question") or "")]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=first + i, column=c, value=v)
            if 7 <= c <= 12 and isinstance(v, float):
                cell.number_format = "0.00%"
    ws.auto_filter.ref = f"A{raw_title + 1}:{get_column_letter(len(RAW))}{last}"
    ws.column_dimensions[get_column_letter(len(RAW))].width = 70


def verdict_banner(ws, r, tables, ncols=9):
    t1, t2, t3, t4 = tables
    wins = {"A": [], "B": []}
    for row in t1[2]:
        w = better(row[2], row[3])
        if w in wins:
            wins[w].append(row[0])
    lines = []
    for s in ("B", "A"):
        if wins[s]:
            lines.append(f"{SYSTEM[s]} scores higher on {len(wins[s])} of 3 metrics: "
                         + ", ".join(wins[s]) + ".")
    overall = "B" if len(wins["B"]) > len(wins["A"]) else "A" if len(wins["A"]) > len(wins["B"]) else None
    # the title states only what the tests support: higher, and where significantly so
    sig = [row[0] for row in t1[2] if row[5] != "—" and float(row[5]) < 0.05
           and better(row[2], row[3]) == overall]
    if not overall:
        title = "Result: neither system is ahead on the RAGAS metrics"
    elif not sig:
        title = (f"Result: {SYSTEM[overall]} scores slightly higher on {len(wins[overall])} of 3 "
                 "metrics, but no difference is significant in the Wilcoxon test")
    elif len(sig) == len(wins[overall]):
        title = (f"Result: {SYSTEM[overall]} scores higher on {len(wins[overall])} of 3 metrics, "
                 "all significant in the Wilcoxon test (p < 0.05)")
    else:
        title = (f"Result: {SYSTEM[overall]} scores slightly higher on {len(wins[overall])} of 3 "
                 f"metrics; Wilcoxon-significant only for {', '.join(s.lower() for s in sig)} (p < 0.05)")
    detail = []
    for row in t1[2]:
        detail.append(f"• {row[0]}: A {row[2].split(' ±')[0]} vs B {row[3].split(' ±')[0]} "
                      f"(Δ {row[4]}, Wilcoxon p = {row[5]})")
    detail.append("• Questions won (Table 2, sign test): " + "; ".join(
        f"{row[0]} A {row[2].split(' (')[0]} vs B {row[3].split(' (')[0]} (p = {row[5]})"
        for row in t2[2]))
    box = [(title, Font(bold=True, size=14, color="006100"))] + \
          [(x, Font(size=11)) for x in lines + detail] + \
          [("Green cell = the better of the two systems in that row (or column); higher is better "
            "for every metric.", NOTE)]
    top = r
    for text, font in box:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = font
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="EAF6EC")
        ws.row_dimensions[r].height = 24 if r == top else 17
        r += 1
    for row in ws.iter_rows(min_row=top, max_row=r - 1, max_col=ncols):
        for cell in row:
            cell.border = Border(top=THIN if cell.row == top else None,
                                 bottom=THIN if cell.row == r - 1 else None,
                                 left=THIN if cell.column == 1 else None,
                                 right=THIN if cell.column == ncols else None)
    return r + 2


def write_mean_sheet(wb, rounds, combined):
    k = len(rounds)
    ws = wb.create_sheet(f"Mean of {k} Runs")
    sheet_header(ws, f"Answer quality with RAGAS metrics — mean of {k} runs",
                 f"batch id: {FINAL} · {len(combined)} test cases scored in every run")
    tables = build(combined)
    r = verdict_banner(ws, 5, tables)

    # headline per run: rows come in A/B pairs, greens per column, winner by mean
    title, head, body = per_run_table(rounds)
    start = r
    r = write_table(ws, r, title, head + ["Better system (mean)"], [row + [""] for row in body])
    first_row = start + 2
    for i in range(0, len(body), 2):
        ra, rb = first_row + i, first_row + i + 1
        for c in range(2, k + 3):
            w = better(ws.cell(row=ra, column=c).value, ws.cell(row=rb, column=c).value)
            if w in ("A", "B"):
                ws.cell(row=ra if w == "A" else rb, column=c).fill = WIN_FILL
        w = better(ws.cell(row=ra, column=k + 2).value, ws.cell(row=rb, column=k + 2).value)
        wc = len(head) + 1
        ws.merge_cells(start_row=ra, start_column=wc, end_row=rb, end_column=wc)
        cell = ws.cell(row=ra, column=wc, value=SYSTEM.get(w, "—"))
        cell.font = Font(bold=True, color="006100")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=r - 1, column=1, value=(
        f"The tables below use, for every answer, the mean of its {k} scores. Tests are computed on "
        "these combined values.")).font = NOTE
    r += 1

    t1, t2, t3, t4 = tables
    r = write_table(ws, r, *t1, pairs=[(3, 4)], winner_head="Higher score")
    r = write_table(ws, r, *t2, pairs=[(3, 4)], winner_head="More questions")
    for t in (t3, t4):
        body = []
        for row in t[2]:
            ws_ = [better(row[2 + 2 * i], row[3 + 2 * i]) for i in range(3)]
            nb, na = ws_.count("B"), ws_.count("A")
            body.append(row + [f"{SYSTEM['B'] if nb > na else SYSTEM['A'] if na > nb else 'Tie'}"
                               f" · {max(na, nb)} of 3"])
        r = write_table(ws, r, t[0], t[1] + ["Better on more metrics"], body,
                        pairs=[(3, 4), (5, 6), (7, 8)])
        for rr in range(r - 2 - len(body), r - 2):    # write_table returns last row + 3
            ws.cell(row=rr, column=len(t[1]) + 1).font = Font(bold=True, color="006100")
    r = write_table(ws, r, *stability_table(rounds))
    for col in "GH":
        ws.column_dimensions[col].width = 28
    ws.column_dimensions["I"].width = 34          # 'Better on more metrics' (tables 3, 4)


def main():
    rounds = load_rounds(logstore.read_all())
    rounds = {k: v for k, v in rounds.items() if k in ROUNDS}
    if not rounds:
        raise SystemExit("No RAGAS rows. Run: python -m app.ragas_eval --batch "
                         f"{FINAL} --round 1")
    wb = Workbook()
    wb.remove(wb.active)
    for k, recs in rounds.items():
        write_run_sheet(wb, k, recs)
        print(f"run {k}: {len(recs)} test cases")
    if len(rounds) > 1:
        comb = combine(rounds)
        write_mean_sheet(wb, rounds, comb)
        for title, head, body in [per_run_table(rounds)] + build(comb) + [stability_table(rounds)]:
            print(f"\n### {title}\n\n{md_table(head, body)}")
    else:
        for title, head, body in build(next(iter(rounds.values()))):
            print(f"\n### {title}\n\n{md_table(head, body)}")
    wb.save(OUT)
    print("\nwritten:", OUT)


if __name__ == "__main__":
    main()

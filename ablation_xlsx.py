"""Write the ablation tables to an Excel workbook, one sheet per ablation.

Each condition gets a mean column and a standard-deviation column, with real
numbers (not text), so the cells can be charted or reformatted.

    python ablation_xlsx.py            -> ..\\ablation-tables.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import logstore, metrics

BASE = "batch-bb633707"
OUT = Path(__file__).resolve().parent.parent / "ablation-tables.xlsx"

# label, key, number format, from per-type block?
ROWS_A = [
    ("Recall@k",        "recall",     "0%",    None),
    ("Precision@k",     "precision",  "0%",    None),
    ("F1",              "f1",         "0%",    None),
    ("MRR",             "mrr",        "0.00",  None),
    ("Pages retrieved", "pages",      "0.0",   None),
    ("Latency (s)",     "latency_ms", "0.0",   None),
    ("Tokens",          "tokens",     "#,##0", None),
    (None, None, None, None),
    ("Recall@k — numeric",   "recall", "0%", "numeric"),
    ("F1 — numeric",         "f1",     "0%", "numeric"),
    ("Recall@k — multi-hop", "recall", "0%", "multi-hop"),
    ("F1 — multi-hop",       "f1",     "0%", "multi-hop"),
]

ROWS_B = [
    ("Recall",            "recall",          "0%",    None),
    ("Precision",         "precision",       "0%",    None),
    ("F1",                "f1",              "0%",    None),
    ("Reads to gold",     "reads_to_gold",   "0.0",   None),
    ("Pages read",        "pages",           "0.0",   None),
    ("Latency (s)",       "latency_ms",      "0.0",   None),
    ("Tokens",            "tokens",          "#,##0", None),
    ("Tool calls",        "effort",          "0.0",   None),
    ("Budget exhausted",  "hit_limit_rate",  "0%",    None),
    (None, None, None, None),
    ("Recall — numeric",   "recall", "0%", "numeric"),
    ("F1 — numeric",       "f1",     "0%", "numeric"),
    ("Recall — multi-hop", "recall", "0%", "multi-hop"),
    ("F1 — multi-hop",     "f1",     "0%", "multi-hop"),
]

GROUPS = [
    ("top-k", "Retrieval depth (top-k)", "A", ROWS_A, "Pipeline A · Vector-Based RAG",
     [("k = 3", "batch-6d5f93d3"), ("k = 5 (baseline)", BASE), ("k = 10", "batch-59c4d149")]),
    ("Embedding model", "Embedding model", "A", ROWS_A, "Pipeline A · Vector-Based RAG",
     [("BGE-M3 (baseline)", BASE), ("multilingual-E5-base", "batch-e9f57386")]),
    ("Chunking", "Chunking strategy", "A", ROWS_A, "Pipeline A · Vector-Based RAG",
     [("clause-aware (baseline)", BASE), ("character windows", "batch-c9000fd0")]),
    ("Step budget", "Step budget", "B", ROWS_B, "Pipeline B · Reasoning-Based Retrieval",
     [("2 steps", "batch-6a6feffd"), ("3 steps", "batch-4bc10bd9"),
      ("8 steps (baseline)", BASE), ("12 steps", "batch-038c5e42")]),
]

HEAD_FILL = PatternFill("solid", fgColor="3B5B33")
SUB_FILL = PatternFill("solid", fgColor="EDF1E7")
BASE_FILL = PatternFill("solid", fgColor="F1F5F2")
THIN = Side(style="thin", color="D0D5CC")


def stat(block, key, type_key):
    """Return (mean, sd); hit_limit_rate is a bare float with no spread."""
    src = block[1].get(type_key, {}) if type_key else block[0]
    v = src.get(key)
    if v is None:
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    return v.get("mean"), v.get("std")


FINAL = "batch-46ded84d"

# label, key, format, A?, B?
FINAL_ROWS = [
    ("Recall",           "recall",        "0%",    True,  True),
    ("Precision",        "precision",     "0%",    True,  True),
    ("F1",               "f1",            "0%",    True,  True),
    ("MRR",              "mrr",           "0.00",  True,  False),
    ("Reads to gold",    "reads_to_gold", "0.0",   False, True),
    ("Pages",            "pages",         "0.0",   True,  True),
    ("Latency (s)",      "latency_ms",    "0.0",   True,  True),
    ("Tokens",           "tokens",        "#,##0", True,  True),
    ("Calls / tool calls", "effort",      "0.0",   True,  True),
    ("Budget exhausted", "hit_limit_rate", "0%",   False, True),
]


def final_sheet(wb, rows):
    sel = [r for r in rows if r.get("batch_id") == FINAL]
    m = metrics.summarize(sel)
    A, B = m["overall"]["A"], m["overall"]["B"]
    errs = sum(1 for r in sel if r.get("error"))

    ws = wb.create_sheet("Final run", 0)
    ws["A1"] = "Final run — Vector-Based RAG vs Reasoning-Based Retrieval"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"89 questions x 3 repeats = 267 test cases, all pairs complete "
                f"({A['recall']['n']} on gold-labelled questions, 24 unanswerable controls). "
                f"{errs} failed runs.")
    ws["A2"].font = Font(size=9, color="5B646E")
    ws["A3"] = f"batch id: {FINAL} · top_k = 5 · step budget = 12 · BGE-M3 · clause-aware chunking"
    ws["A3"].font = Font(size=9, color="5B646E")

    top, sub = 5, 6
    ws.cell(row=top, column=1, value="Metric")
    ws.merge_cells(start_row=top, start_column=1, end_row=sub, end_column=1)
    for i, label in enumerate(("A · Vector-Based RAG", "B · Reasoning-Based Retrieval")):
        c0 = 2 + i * 2
        ws.cell(row=top, column=c0, value=label)
        ws.merge_cells(start_row=top, start_column=c0, end_row=top, end_column=c0 + 1)
        ws.cell(row=sub, column=c0, value="mean")
        ws.cell(row=sub, column=c0 + 1, value="± sd")
    for r in (top, sub):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(bold=True, color="FFFFFF" if r == top else "1F2328", size=10)
            cell.fill = HEAD_FILL if r == top else SUB_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    r = sub + 1
    for label, key, fmt, in_a, in_b in FINAL_ROWS:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True, size=10)
        for i, (blk, keep) in enumerate(((A, in_a), (B, in_b))):
            mean = sd = None
            if keep:
                mean, sd = stat((blk, {}), key, None)
                if label.startswith("Latency") and mean is not None:
                    mean, sd = mean / 1000, (sd / 1000 if sd else None)
            for j, val in enumerate((mean, sd)):
                cell = ws.cell(row=r, column=2 + i * 2 + j, value=val if keep else "—")
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="center")
        r += 1

    ov = m["overall"]["overlap"]
    ws.cell(row=r, column=1, value="Page overlap (Jaccard)").font = Font(bold=True, size=10)
    ws.cell(row=r, column=2, value=ov.get("mean")).number_format = "0.00"
    ws.cell(row=r, column=3, value=ov.get("std")).number_format = "0.00"
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=4, value="computed per test case, over both evidence sets")
    ws.cell(row=r, column=4).font = Font(size=9, italic=True, color="5B646E")
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)

    for row in ws.iter_rows(min_row=top, max_row=r, max_col=5):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    # per question type
    t0 = r + 3
    ws.cell(row=t0 - 1, column=1, value="By question type").font = Font(bold=True, size=11)
    heads = ["Question type", "n", "A recall", "B recall", "A precision", "B precision",
             "A F1", "B F1"]
    for c, h in enumerate(heads, start=1):
        cell = ws.cell(row=t0, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    rr = t0 + 1
    for t in sorted(m["per_type"], key=lambda x: -x["n"]):
        a, b = t["A"], t["B"]
        ws.cell(row=rr, column=1, value=t["type"])
        ws.cell(row=rr, column=2, value=t["n"])
        for c, st in ((3, a.get("recall")), (4, b.get("recall")), (5, a.get("precision")),
                      (6, b.get("precision")), (7, a.get("f1")), (8, b.get("f1"))):
            v = st.get("mean") if st else None
            cell = ws.cell(row=rr, column=c, value=v if v is not None else "—")
            cell.number_format = "0%"
            cell.alignment = Alignment(horizontal="center")
        rr += 1
    for row in ws.iter_rows(min_row=t0, max_row=rr - 1, max_col=8):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    ws.column_dimensions["A"].width = 24
    for c in range(2, 9):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.cell(row=rr + 1, column=1, value=(
        "Page-level metrics against the gold pages; unanswerable controls carry no gold and are "
        "excluded from recall, precision and F1. MRR applies to A's ranked list, reads-to-gold to "
        "B's read order (averaged over the runs that reached the gold page). 'Budget exhausted' is "
        "the share of B runs cut off by the step budget. Overlap is the Jaccard index of the two "
        "evidence sets.")).font = Font(size=9, italic=True, color="5B646E")


def main():
    rows = logstore.read_all()
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, title, phase, row_spec, pipeline, cols in GROUPS:
        qids = None
        for _, bid in cols:
            s = {r.get("question_id") for r in rows if r.get("batch_id") == bid}
            qids = s if qids is None else (qids & s)
        blocks = []
        for _, bid in cols:
            sel = [r for r in rows if r.get("batch_id") == bid and r.get("question_id") in qids]
            m = metrics.summarize(sel)
            blocks.append((m["overall"][phase], {t["type"]: t[phase] for t in m["per_type"]}))

        ws = wb.create_sheet(sheet_name)
        ws["A1"] = f"Ablation — {title}"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = (f"{pipeline}. {len(qids)} questions (numeric + multi-hop) x 3 repeats. "
                    f"The other pipeline is unaffected by this setting and was not re-run.")
        ws["A2"].font = Font(size=9, color="5B646E")
        ws["A3"] = "batch ids: " + " · ".join(f"{label} = {bid}" for label, bid in cols)
        ws["A3"].font = Font(size=9, color="5B646E")

        top, sub = 5, 6
        ws.cell(row=top, column=1, value="Metric")
        ws.merge_cells(start_row=top, start_column=1, end_row=sub, end_column=1)
        for i, (label, _) in enumerate(cols):
            c0 = 2 + i * 2
            ws.cell(row=top, column=c0, value=label)
            ws.merge_cells(start_row=top, start_column=c0, end_row=top, end_column=c0 + 1)
            ws.cell(row=sub, column=c0, value="mean")
            ws.cell(row=sub, column=c0 + 1, value="± sd")
        for r in (top, sub):
            for c in range(1, 2 + 2 * len(cols)):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(bold=True, color="FFFFFF" if r == top else "1F2328", size=10)
                cell.fill = HEAD_FILL if r == top else SUB_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        r = sub + 1
        for label, key, fmt, type_key in row_spec:
            if label is None:
                r += 1
                continue
            ws.cell(row=r, column=1, value=label).font = Font(bold="—" not in label, size=10)
            for i, blk in enumerate(blocks):
                mean, sd = stat(blk, key, type_key)
                if label.startswith("Latency") and mean is not None:
                    mean, sd = mean / 1000, (sd / 1000 if sd else None)
                for j, val in enumerate((mean, sd)):
                    cell = ws.cell(row=r, column=2 + i * 2 + j)
                    cell.value = val
                    cell.number_format = fmt
                    cell.alignment = Alignment(horizontal="center")
                    if cols[i][1] == BASE:
                        cell.fill = BASE_FILL
            r += 1

        for row in ws.iter_rows(min_row=top, max_row=r - 1, max_col=1 + 2 * len(cols)):
            for cell in row:
                cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

        ws.column_dimensions["A"].width = 24
        for c in range(2, 2 + 2 * len(cols)):
            ws.column_dimensions[get_column_letter(c)].width = 11
        ws.freeze_panes = "B7"

        note = ("Page-level metrics against the gold pages; precision's denominator is the number "
                "of distinct pages, not k. Means over runs, standard deviation across repeats. "
                "The baseline column is shaded.")
        if phase == "B":
            note += (" 'Budget exhausted' is the share of runs in which the agent was still working "
                     "when its step budget ran out and had to answer from partial evidence. "
                     "'Reads to gold' is the position of the first gold page in the read order, "
                     "averaged over the runs that reached it at all.")
        ws.cell(row=r + 1, column=1, value=note).font = Font(size=9, italic=True, color="5B646E")

    final_sheet(wb, rows)
    wb.save(OUT)
    print("written:", OUT)
    for name in wb.sheetnames:
        print("  sheet:", name)


if __name__ == "__main__":
    main()
